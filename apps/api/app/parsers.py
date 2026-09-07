from __future__ import annotations

import io


def extract_document_text(data: bytes, extension: str) -> str:
    extension = extension.lower().lstrip('.')
    if extension in {'txt', 'md'}:
        return data.decode('utf-8', errors='replace')
    try:
        if extension == 'pdf':
            import fitz
            with fitz.open(stream=data, filetype='pdf') as document:
                return '\n\n'.join(page.get_text('text') for page in document)
        if extension == 'docx':
            from docx import Document
            return '\n\n'.join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
        if extension == 'pptx':
            from pptx import Presentation
            return '\n\n'.join(shape.text for slide in Presentation(io.BytesIO(data)).slides for shape in slide.shapes if hasattr(shape, 'text'))
    except Exception as error:
        raise ValueError('DOCUMENT_PARSE_FAILED') from error
    raise ValueError('DOCUMENT_UNSUPPORTED')
